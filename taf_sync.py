"""taf_sync.py - Apollo-IMS-TAF-2.7-Export in die Output-Dateien der Import-GUI uebernehmen.

Liest einen Exportordner (0757\\*.dat mit festen Satzbreiten, UTF-8, CRLF sowie
0757_PIC.zip / 0757_DOC.zip) und schreibt die Saetze in die Excel-Dateien des
GUI-Output-Ordners, damit der GUI-Datenstamm dem Apollo-Stand entspricht.

Merge-Regel: je (Artikel, Zieldatei) gewinnt die GUI, wenn sie fuer den Artikel
in dieser Datei schon Zeilen hat (die GUI ist dort neuer). Dann werden die
IMS-Zeilen fuer diesen Artikel verworfen und als Konflikt mit beiden Staenden
in den Bericht geschrieben. Sonst werden die IMS-Zeilen angehaengt. Vorhandene
GUI-Zeilen werden nicht veraendert - einzige Ausnahme: leere Text-Modul-IDs
werden gefuellt (repair_empty_module_ids, im Bericht unter Hinweise).

Aufruf:
    python -I -X utf8 taf_sync.py <Exportordner> <Output-Ordner> [--no-media]
        [--media-limit N] [--report <xlsx>]

Das Modul aendert apollo_import_gui.py nicht; es benutzt nur deren Konstanten
und Lese-/Schreibfunktionen (Dateinamen, Header, Workbook-Zugriff, Medien-
Namensregeln, IdRegistry, Kataloge).
"""
from __future__ import annotations

import argparse
import re
import sys
import time
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

REPO_DIR = Path(__file__).resolve().parent
if str(REPO_DIR) not in sys.path:
    sys.path.insert(0, str(REPO_DIR))

import apollo_import_gui as m  # noqa: E402

DATA_SUPPLIER_NO = "0757"
# 209 (GTIN) wird nur gezaehlt: die GUI fuehrt bewusst keine EAN (Entscheidung 26.10.2026).
# Dateien, bei denen die Zeilenreihenfolge fuer den Konfliktvergleich keine Rolle spielt.
ORDER_INSENSITIVE_FILES = {m.GENART_FILE[0], m.VIDEO_FILE[0], m.WEB_LINK_FILE[0]}
LOGO_SUBFOLDER = "_logo"
MEDIA_SOURCE_PREFIX = "taf:"
# TAF-Sprachnummer -> Schluessel der TranslationSet-Felder der App.
LANG_NO_TO_KEY = {"255": "uni", "001": "de", "004": "en", "018": "cz", "006": "fr", "007": "it", "009": "nl"}
# Beleg-Formattyp (DT 014) -> Dateiendung im Zip (gross) und in der GUI (klein).
DOC_TYPE_ZIP_SUFFIX = {"02": ".PDF", "03": ".JPG", "06": ".PNG"}
DOC_TYPE_GUI_SUFFIX = {"02": ".pdf", "03": ".jpg", "06": ".png"}
IMAGE_DOC_TYPES = ("03", "06")
DOCUMENT_DOC_TYPES = ("02",)
URL_DOC_TYPE = "04"
LOGO_CONTENT_TYPE = "002"
CONFLICT_TEXT_MAX_LENGTH = 2000

# Satzaufbau je Tabelle: (Satzlaenge, [(Feldname, Startposition 1-basiert, Breite)]).
# Nur die Felder, die ein GUI-Ziel haben; alle anderen sind laut Gesamtkarte konstant.
TABLE_LAYOUTS: dict[str, tuple[int, list[tuple[str, int, int]]]] = {
    "200": (53, [("ArtNo", 1, 22), ("TermNo", 30, 9), ("Accessory", 42, 1)]),
    "030": (102, [("TermNo", 30, 9), ("LangNo", 39, 3), ("Term", 42, 60)]),
    "035": (103, [("TxtModNo", 30, 6), ("LangNo", 36, 3), ("SeqNo", 40, 3), ("Text", 43, 60)]),
    # 206: Reserved 30-34, CountryCode 35-37, SortNo 38-39, InfType 40-42, FirstPage 43,
    # TxtModNo 44-49, Exclude 50, DeleteFlag 51 (am Export geprueft; die Gesamtkarte
    # nennt InfType 41/TxtModNo 45, das laesst Position 40 aus).
    "206": (51, [("ArtNo", 1, 22), ("SortNo", 38, 2), ("InfType", 40, 3), ("TxtModNo", 44, 6)]),
    "207": (75, [("ArtNo", 1, 22), ("TradeNo", 33, 35), ("SortNo", 70, 5)]),
    "203": (71, [("ArtNo", 1, 22), ("ManNo", 30, 6), ("RefNo", 39, 22), ("SortNo", 62, 5)]),
    "209": (48, [("ArtNo", 1, 22), ("GTIN", 33, 14)]),
    "210": (67, [("ArtNo", 1, 22), ("SortNo", 38, 3), ("CritNo", 41, 4), ("CritVal", 45, 20)]),
    "211": (35, [("ArtNo", 1, 22), ("GenArtNo", 30, 5)]),
    "212": (54, [("ArtNo", 1, 22), ("ArtStat", 43, 3), ("StatusDat", 46, 8)]),
    "222": (
        83,
        [
            ("ArtNo", 1, 22),
            ("SeqNo", 37, 3),
            ("SortNo", 40, 3),
            ("AccArtNo", 43, 22),
            ("Quantity", 65, 3),
            ("AccGenArtNo", 68, 5),
        ],
    ),
    "205": (63, [("ArtNo", 1, 22), ("PartGenArtNo", 30, 5), ("SeqNo", 35, 3), ("PartNo", 38, 22), ("Quantity", 60, 3)]),
    "231": (
        400,
        [
            ("DocNo", 30, 9),
            ("LangNo", 39, 3),
            ("DocName", 42, 30),
            ("DocContentType", 72, 3),
            ("DocTermNorm", 75, 3),
            ("DocType", 89, 2),
            ("URL", 101, 300),
        ],
    ),
    "232": (47, [("ArtNo", 1, 22), ("SortNo", 30, 2), ("DocNo", 36, 9), ("DocType", 45, 2)]),
    "299": (61, [("ArtNo", 1, 22), ("ArticleGUID", 30, 32)]),
    "400": (56, [("ArtNo", 1, 22), ("GenArtNo", 30, 5), ("LnkTargetType", 35, 3), ("LnkTargetNo", 38, 9)]),
}
# Felder, bei denen nur rechts getrimmt wird (innere/fuehrende Blanks gehoeren zum Inhalt).
RSTRIP_ONLY_FIELDS = {"ArtNo", "Term", "Text", "AccArtNo", "PartNo", "DocName", "URL"}


# ---------------------------------------------------------------------------
# Datenobjekte des Exports
# ---------------------------------------------------------------------------


@dataclass(slots=True)
class TafArticle:
    art_no: str
    term_no: str
    accessory: bool


@dataclass(slots=True)
class TafInfoText:
    art_no: str
    sort_no: str
    inf_type: str
    txt_mod_no: str


@dataclass(slots=True)
class TafTradeNumber:
    art_no: str
    trade_no: str
    sort_no: int


@dataclass(slots=True)
class TafReferenceNumber:
    art_no: str
    man_no: str
    ref_no: str
    sort_no: int


@dataclass(slots=True)
class TafGtin:
    art_no: str
    gtin: str


@dataclass(slots=True)
class TafCriterion:
    art_no: str
    sort_no: int
    crit_no: str
    crit_val: str


@dataclass(slots=True)
class TafGenArt:
    art_no: str
    genart_no: str


@dataclass(slots=True)
class TafStatus:
    art_no: str
    art_stat: str
    status_date: str


@dataclass(slots=True)
class TafAccessory:
    art_no: str
    seq_no: str
    sort_no: int
    acc_art_no: str
    quantity: str
    acc_genart_no: str


@dataclass(slots=True)
class TafPart:
    art_no: str
    part_genart_no: str
    seq_no: str
    part_no: str
    quantity: str


@dataclass(slots=True)
class TafDocument:
    doc_no: str
    doc_type: str
    lang_no: str
    doc_name: str
    content_type: str
    term_norm: str
    url: str


@dataclass(slots=True)
class TafAllocation:
    art_no: str
    sort_no: int
    doc_no: str
    doc_type: str


@dataclass(slots=True)
class TafGuid:
    art_no: str
    guid: str


@dataclass(slots=True)
class TafLinkage:
    art_no: str
    genart_no: str
    target_type: str
    target_no: str


@dataclass
class TafExport:
    export_dir: Path
    data_dir: Path
    pic_zip: Path | None = None
    doc_zip: Path | None = None
    articles: list[TafArticle] = field(default_factory=list)
    # (TermNo, LangNo) -> Kurzbezeichnung
    terms: dict[tuple[str, str], str] = field(default_factory=dict)
    # TxtModNo -> LangNo -> [(SeqNo, Zeile)] in Dateireihenfolge
    text_modules: dict[str, dict[str, list[tuple[int, str]]]] = field(default_factory=dict)
    info_texts: list[TafInfoText] = field(default_factory=list)
    trade_numbers: list[TafTradeNumber] = field(default_factory=list)
    reference_numbers: list[TafReferenceNumber] = field(default_factory=list)
    gtins: list[TafGtin] = field(default_factory=list)
    criteria: list[TafCriterion] = field(default_factory=list)
    genarts: list[TafGenArt] = field(default_factory=list)
    statuses: list[TafStatus] = field(default_factory=list)
    accessory_lists: list[TafAccessory] = field(default_factory=list)
    parts_lists: list[TafPart] = field(default_factory=list)
    # (DocNo, DocType) -> 231-Satz (Sprache 255 bzw. erster Satz)
    documents: dict[tuple[str, str], TafDocument] = field(default_factory=dict)
    allocations: list[TafAllocation] = field(default_factory=list)
    guids: list[TafGuid] = field(default_factory=list)
    linkages: list[TafLinkage] = field(default_factory=list)
    record_counts: dict[str, int] = field(default_factory=dict)

    @property
    def article_numbers(self) -> list[str]:
        return [article.art_no for article in self.articles]


# ---------------------------------------------------------------------------
# Parser
# ---------------------------------------------------------------------------


def strip_leading_zeros(value: str) -> str:
    """'015040' -> '15040', '000' -> '0', '' -> ''."""
    value = value.strip()
    if not value:
        return ""
    stripped = value.lstrip("0")
    return stripped or "0"


def gtin_to_ean(value: str) -> str:
    """GTIN-14 mit fuehrender 0 auf EAN-13 kuerzen; andere Laengen unveraendert."""
    digits = value.strip()
    if len(digits) == 14 and digits.startswith("0"):
        return digits[1:]
    return digits


def split_numeric_interval(value: str) -> tuple[str, str]:
    """'10,3-12,8' -> ('10,3', '12,8'); '-5--40' -> ('-5', '-40'); '7,5' -> ('7,5', '')."""
    text = value.strip()
    match = re.fullmatch(r"(-?\d+(?:[.,]\d+)?)-(-?\d+(?:[.,]\d+)?)", text)
    if match:
        return match.group(1), match.group(2)
    return text, ""


def parse_fixed_records(lines: list[str], table: str, *, source: str = "") -> list[dict[str, str]]:
    """Schneidet Zeilen exakt nach den Positionen der Tabelle; prueft Satzlaenge und TableNo."""
    record_length, fields = TABLE_LAYOUTS[table]
    records: list[dict[str, str]] = []
    for line_no, line in enumerate(lines, start=1):
        if line == "":
            continue
        if len(line) != record_length:
            raise ValueError(
                f"{source or table}.dat Zeile {line_no}: Satzlaenge {len(line)} statt {record_length}: {line[:40]!r}"
            )
        if line[26:29] != table:
            raise ValueError(f"{source or table}.dat Zeile {line_no}: TableNo {line[26:29]!r} statt {table!r}")
        record: dict[str, str] = {}
        for name, start, width in fields:
            raw = line[start - 1 : start - 1 + width]
            record[name] = raw.rstrip() if name in RSTRIP_ONLY_FIELDS else raw.strip()
        records.append(record)
    return records


def _read_table_lines(data_dir: Path, table: str) -> list[str]:
    path = data_dir / f"{table}.dat"
    if not path.exists():
        raise FileNotFoundError(f"TAF-Tabelle fehlt: {path}")
    # Bytes dekodieren statt read_text: die Zeilenenden (CRLF) muessen erhalten bleiben.
    text = path.read_bytes().decode("utf-8")
    if not text:
        return []
    if text.startswith("﻿"):
        text = text[1:]
    if "\r\n" not in text and "\n" in text:
        raise ValueError(f"{path.name}: Zeilenende ist nicht CRLF")
    lines = text.split("\r\n")
    if lines and lines[-1] == "":
        lines.pop()
    for line in lines:
        if "\n" in line or "\r" in line:
            raise ValueError(f"{path.name}: Zeilenumbruch innerhalb eines Satzes")
    return lines


def _read_table(data_dir: Path, table: str) -> list[dict[str, str]]:
    return parse_fixed_records(_read_table_lines(data_dir, table), table, source=table)


def _to_int(value: str) -> int:
    value = value.strip()
    return int(value) if value.isdigit() else 0


def read_taf_export(export_dir: Path | str) -> TafExport:
    """Liest einen Exportordner (<Ordner>\\0757\\*.dat, 0757_PIC.zip, 0757_DOC.zip)."""
    export_dir = Path(export_dir)
    data_dir = export_dir / DATA_SUPPLIER_NO
    if not data_dir.is_dir():
        if (export_dir / "200.dat").exists():
            data_dir = export_dir
        else:
            raise FileNotFoundError(f"Kein Datenordner {DATA_SUPPLIER_NO} in {export_dir}")
    export = TafExport(export_dir=export_dir, data_dir=data_dir)
    pic_zip = export_dir / f"{DATA_SUPPLIER_NO}_PIC.zip"
    doc_zip = export_dir / f"{DATA_SUPPLIER_NO}_DOC.zip"
    export.pic_zip = pic_zip if pic_zip.exists() else None
    export.doc_zip = doc_zip if doc_zip.exists() else None

    def table(name: str) -> list[dict[str, str]]:
        records = _read_table(data_dir, name)
        export.record_counts[name] = len(records)
        return records

    for record in table("200"):
        export.articles.append(TafArticle(art_no=record["ArtNo"], term_no=record["TermNo"], accessory=record["Accessory"] == "1"))
    for record in table("030"):
        export.terms[(record["TermNo"], record["LangNo"])] = record["Term"]
    for record in table("035"):
        export.text_modules.setdefault(record["TxtModNo"], {}).setdefault(record["LangNo"], []).append(
            (_to_int(record["SeqNo"]), record["Text"])
        )
    for record in table("206"):
        export.info_texts.append(
            TafInfoText(art_no=record["ArtNo"], sort_no=record["SortNo"], inf_type=record["InfType"], txt_mod_no=record["TxtModNo"])
        )
    for record in table("207"):
        export.trade_numbers.append(TafTradeNumber(art_no=record["ArtNo"], trade_no=record["TradeNo"], sort_no=_to_int(record["SortNo"])))
    for record in table("203"):
        export.reference_numbers.append(
            TafReferenceNumber(art_no=record["ArtNo"], man_no=record["ManNo"], ref_no=record["RefNo"], sort_no=_to_int(record["SortNo"]))
        )
    for record in table("209"):
        export.gtins.append(TafGtin(art_no=record["ArtNo"], gtin=record["GTIN"]))
    for record in table("210"):
        export.criteria.append(
            TafCriterion(art_no=record["ArtNo"], sort_no=_to_int(record["SortNo"]), crit_no=record["CritNo"], crit_val=record["CritVal"])
        )
    for record in table("211"):
        export.genarts.append(TafGenArt(art_no=record["ArtNo"], genart_no=record["GenArtNo"]))
    for record in table("212"):
        export.statuses.append(TafStatus(art_no=record["ArtNo"], art_stat=record["ArtStat"], status_date=record["StatusDat"]))
    for record in table("222"):
        export.accessory_lists.append(
            TafAccessory(
                art_no=record["ArtNo"],
                seq_no=record["SeqNo"],
                sort_no=_to_int(record["SortNo"]),
                acc_art_no=record["AccArtNo"],
                quantity=record["Quantity"],
                acc_genart_no=record["AccGenArtNo"],
            )
        )
    for record in table("205"):
        export.parts_lists.append(
            TafPart(
                art_no=record["ArtNo"],
                part_genart_no=record["PartGenArtNo"],
                seq_no=record["SeqNo"],
                part_no=record["PartNo"],
                quantity=record["Quantity"],
            )
        )
    for record in table("231"):
        document = TafDocument(
            doc_no=record["DocNo"],
            doc_type=record["DocType"],
            lang_no=record["LangNo"],
            doc_name=record["DocName"],
            content_type=record["DocContentType"],
            term_norm=record["DocTermNorm"],
            url=record["URL"],
        )
        key = (document.doc_no, document.doc_type)
        # Je DocNo/DocType muss ein 255-Satz existieren; der gewinnt.
        if key not in export.documents or document.lang_no == "255":
            export.documents[key] = document
    for record in table("232"):
        export.allocations.append(
            TafAllocation(art_no=record["ArtNo"], sort_no=_to_int(record["SortNo"]), doc_no=record["DocNo"], doc_type=record["DocType"])
        )
    for record in table("299"):
        export.guids.append(TafGuid(art_no=record["ArtNo"], guid=record["ArticleGUID"]))
    for record in table("400"):
        export.linkages.append(
            TafLinkage(
                art_no=record["ArtNo"],
                genart_no=record["GenArtNo"],
                target_type=record["LnkTargetType"],
                target_no=record["LnkTargetNo"],
            )
        )
    return export


TEXT_MODULE_LINE_WIDTH = 60


def join_text_module_lines(lines: list[tuple[int, str]]) -> str:
    """Textbaustein-Zeilen in SeqNo-Reihenfolge zu Absaetzen zusammensetzen.

    IMS bricht Absaetze beim Export auf 60 Zeichen um (am letzten Blank innerhalb
    der ersten 61 Zeichen, dieses Blank entfaellt). Der Umbruch wird hier
    rueckgaengig gemacht: Haette das erste Wort der Folgezeile noch in die
    Zeile gepasst (Laenge + Blank + Wort <= 60), war der Umbruch gewollt - sonst
    ist er ein Exportumbruch und die Zeilen gehoeren zu einem Absatz. Satz mit
    60 Blanks = Leerzeile (Absatztrenner). Die Zeilen sind beim Parsen rechts
    getrimmt; Leerzeilen am Modulende entfallen.
    """
    ordered = [text for _seq, text in sorted(lines, key=lambda item: item[0])]
    while ordered and ordered[-1] == "":
        ordered.pop()
    paragraphs: list[str] = []
    last_line = ""  # letzte physische Exportzeile (nicht der zusammengefuegte Absatz)
    for text in ordered:
        first_word = text.split(" ", 1)[0]
        if text and paragraphs and last_line and len(last_line) + 1 + len(first_word) > TEXT_MODULE_LINE_WIDTH:
            paragraphs[-1] = paragraphs[-1] + " " + text
        else:
            paragraphs.append(text)
        last_line = text
    return "\n".join(paragraphs)


def short_text_translations(export: TafExport, term_no: str) -> dict[str, str]:
    """Kurzbezeichnung je Sprachschluessel (uni/de/en/cz/fr/it/nl) fuer eine TermNo."""
    values: dict[str, str] = {}
    for lang_no, key in LANG_NO_TO_KEY.items():
        values[key] = export.terms.get((term_no, lang_no), "")
    return values


def long_text_translations(export: TafExport, txt_mod_no: str) -> dict[str, str]:
    module = export.text_modules.get(txt_mod_no, {})
    values: dict[str, str] = {}
    for lang_no, key in LANG_NO_TO_KEY.items():
        values[key] = join_text_module_lines(module.get(lang_no, []))
    return values


def build_text_row(article: str, module_id: str, module_type: str, translations: dict[str, str]) -> list[str]:
    """Zeile im Aufbau von SHORT_TEXT_HEADERS (ohne Zeitstempel)."""
    row = [article, module_id, module_type, "1"]
    for language_id, key in m.EXPORT_LANGUAGE_LAYOUT:
        row.append(language_id)
        row.append(translations.get(key, ""))
    return row


def classify_url(url: str) -> tuple[str, str]:
    """('video', Embed-Link) fuer YouTube, sonst ('web', Link)."""
    value = url.strip()
    if m.extract_youtube_video_id(value) or "youtube" in value.lower():
        return "video", m.normalize_youtube_url_for_embed(value)
    return "web", value


def media_target_dir(output_dir: Path, article: str) -> Path:
    """Artikel-Medienordner wie materialize_media_rows: <output>\\Medien\\<Artikel>, aufgeloest
    (resolve() liefert auf Netzlaufwerken den UNC-Pfad, so entstehen die Pfade der GUI)."""
    return (Path(output_dir) / m.MEDIA_SUBFOLDER / m.safe_folder_name(article)).resolve()


# ---------------------------------------------------------------------------
# Bericht
# ---------------------------------------------------------------------------


@dataclass
class FileStats:
    file_name: str
    ims_rows: int = 0
    appended: int = 0
    kept: int = 0
    conflict_articles: int = 0
    identical_conflicts: int = 0
    discarded: int = 0
    written: bool = False


@dataclass
class Conflict:
    article: str
    file_name: str
    gui_state: str
    ims_state: str
    # True, wenn GUI- und IMS-Zeilen (ohne Zeitstempel) inhaltlich gleich sind.
    identical: bool = False


@dataclass
class Hint:
    category: str
    article: str
    message: str


@dataclass
class TafSyncReport:
    export_dir: str
    output_dir: str
    written_at: str
    files: dict[str, FileStats] = field(default_factory=dict)
    conflicts: list[Conflict] = field(default_factory=list)
    hints: list[Hint] = field(default_factory=list)
    counted_only: dict[str, int] = field(default_factory=dict)
    media_articles: int = 0
    media_files_written: int = 0
    media_files_reused: int = 0
    logo_path: str = ""
    duration_seconds: float = 0.0
    report_path: str = ""

    def stats(self, file_name: str) -> FileStats:
        stats = self.files.get(file_name)
        if stats is None:
            stats = FileStats(file_name=file_name)
            self.files[file_name] = stats
        return stats

    def add_hint(self, category: str, article: str, message: str) -> None:
        self.hints.append(Hint(category=category, article=article, message=message))

    def summary_lines(self) -> list[str]:
        lines = [f"TAF-Sync {self.export_dir} -> {self.output_dir} ({self.written_at})"]
        for stats in self.files.values():
            lines.append(
                f"  {stats.file_name}: IMS {stats.ims_rows}, uebernommen {stats.appended}, GUI behalten {stats.kept}, "
                f"Konflikt-Artikel {stats.conflict_articles} (davon inhaltsgleich {stats.identical_conflicts}), verworfen {stats.discarded}"
                + ("" if stats.written else " (nicht geschrieben)")
            )
        for table, count in self.counted_only.items():
            lines.append(f"  {table}: {count} Saetze nur gezaehlt (kein GUI-Ziel)")
        lines.append(
            f"  Medien: {self.media_articles} Artikel, {self.media_files_written} Dateien neu, {self.media_files_reused} wiederverwendet"
            + (f", Logo {self.logo_path}" if self.logo_path else "")
        )
        lines.append(f"  Konflikte: {len(self.conflicts)}, Hinweise: {len(self.hints)}, Laufzeit {self.duration_seconds:.1f} s")
        if self.report_path:
            lines.append(f"  Bericht: {self.report_path}")
        return lines

    def write_xlsx(self, path: Path | str) -> Path:
        from openpyxl import Workbook

        path = Path(path)
        workbook = Workbook()
        overview = workbook.active
        overview.title = "Uebersicht"
        overview.append(["Export", self.export_dir])
        overview.append(["Output", self.output_dir])
        overview.append(["Zeitstempel", self.written_at])
        overview.append(["Laufzeit [s]", f"{self.duration_seconds:.1f}"])
        overview.append([])
        overview.append(
            ["Datei", "IMS-Zeilen", "uebernommen", "GUI behalten", "Konflikt-Artikel", "davon inhaltsgleich", "verworfen", "geschrieben"]
        )
        for stats in self.files.values():
            overview.append(
                [
                    stats.file_name,
                    stats.ims_rows,
                    stats.appended,
                    stats.kept,
                    stats.conflict_articles,
                    stats.identical_conflicts,
                    stats.discarded,
                    "ja" if stats.written else "nein",
                ]
            )
        for table, count in self.counted_only.items():
            overview.append([f"{table} (nur gezaehlt)", count, 0, 0, 0, 0, count, "nein"])
        overview.append([])
        overview.append(["Medien-Artikel", self.media_articles])
        overview.append(["Mediendateien neu", self.media_files_written])
        overview.append(["Mediendateien wiederverwendet", self.media_files_reused])
        overview.append(["Logo", self.logo_path])
        overview.append(["Konflikte", len(self.conflicts)])
        overview.append(["Hinweise", len(self.hints)])

        conflicts = workbook.create_sheet("Konflikte")
        conflicts.append(["Artikel", "Datei", "inhaltsgleich", "GUI-Stand", "IMS-Stand"])
        for conflict in self.conflicts:
            conflicts.append(
                [conflict.article, conflict.file_name, "ja" if conflict.identical else "nein", conflict.gui_state, conflict.ims_state]
            )

        hints = workbook.create_sheet("Hinweise")
        hints.append(["Kategorie", "Artikel", "Hinweis"])
        for hint in self.hints:
            hints.append([hint.category, hint.article, hint.message])

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        self.report_path = str(path)
        return path


# ---------------------------------------------------------------------------
# Merge
# ---------------------------------------------------------------------------


def _compact_rows(rows: list[list[str]], drop_last: bool) -> str:
    parts = []
    for row in rows:
        cells = row[:-1] if drop_last and row else row
        parts.append(" | ".join(str(cell) for cell in cells))
    text = "; ".join(parts)
    if len(text) > CONFLICT_TEXT_MAX_LENGTH:
        text = text[: CONFLICT_TEXT_MAX_LENGTH - 3] + "..."
    return text


class _Merger:
    """Haelt Kataloge, Zeitstempel und Bericht fuer einen Lauf."""

    def __init__(
        self,
        export: TafExport,
        output_dir: Path,
        *,
        media: bool,
        media_limit: int | None,
        written_at: str,
        genart_catalog_path: Path,
        attribute_catalog_path: Path,
        competitor_catalog_path: Path,
        vehicle_catalog_path: Path | None,
    ) -> None:
        self.export = export
        self.output_dir = output_dir
        self.media = media
        self.media_limit = media_limit
        self.written_at = written_at
        self.report = TafSyncReport(export_dir=str(export.export_dir), output_dir=str(output_dir), written_at=written_at)
        self.key_of: dict[str, str] = {}
        self.articles_by_key: dict[str, list[str]] = {}
        for article in export.articles:
            key = m.normalize_article_number(article.art_no)
            self.key_of[article.art_no] = key
            self.articles_by_key.setdefault(key, []).append(article.art_no)
        for key, spellings in self.articles_by_key.items():
            if len(spellings) > 1:
                self.report.add_hint(
                    "Artikelnummer",
                    spellings[0],
                    "Mehrere IMS-Schreibweisen fallen in der GUI zusammen: " + ", ".join(repr(value) for value in spellings),
                )
        self.genart_names = self._load_genart_names(genart_catalog_path)
        self.attribute_options = self._load_attribute_options(attribute_catalog_path)
        self.manufacturers, self.competitors = self._load_competitor_catalog(competitor_catalog_path)
        self.vehicle_catalog = self._load_vehicle_catalog(vehicle_catalog_path)
        self.missing_genarts: set[str] = set()
        self.missing_criteria: set[str] = set()

    # -- Kataloge -----------------------------------------------------------

    def _load_genart_names(self, path: Path) -> dict[str, str]:
        if not m.path_exists_safe(path):
            self.report.add_hint("Katalog", "", f"GenArt-Katalog nicht gefunden: {path} - Bezeichnungen bleiben leer")
            return {}
        registry = m.GenArtRegistry()
        registry.load_from_workbook(Path(path))
        return {genart_id: option.bezeichnung for genart_id, option in registry.options_by_id.items()}

    def _load_attribute_options(self, path: Path) -> dict[str, m.AttributeOption]:
        if not m.path_exists_safe(path):
            self.report.add_hint("Katalog", "", f"Attribut-Katalog nicht gefunden: {path} - Bezeichnung/Format bleiben leer")
            return {}
        return {option.criteria_id: option for option in m.load_attribute_options(Path(path))}

    def _load_competitor_catalog(self, path: Path) -> tuple[dict[str, m.CompetitorOption], dict[str, m.CompetitorOption]]:
        if not m.path_exists_safe(path):
            self.report.add_hint("Katalog", "", f"KHer-Katalog nicht gefunden: {path} - alle Referenznummern werden OE-Nummern")
            return {}, {}
        manufacturers = {option.competitor_id: option for option in m.load_competitor_options(Path(path), comparison_only=False)}
        competitors = {option.competitor_id: option for option in m.load_competitor_options(Path(path), comparison_only=True)}
        return manufacturers, competitors

    def _load_vehicle_catalog(self, path: Path | None):
        if path is None or not m.path_exists_safe(path):
            return None
        try:
            return m.load_vehicle_catalog(Path(path))
        except Exception as exc:  # Katalog ist nur fuer Hinweise noetig
            self.report.add_hint("Katalog", "", f"KTyp-Katalog nicht ladbar ({exc}) - keine Topmotive-Pruefung")
            return None

    def _ordered(self, records, *keys):
        """Saetze je Artikel in 200-Reihenfolge, darin nach den genannten Attributen
        (SortNo numerisch) - unabhaengig von der Reihenfolge in der Datei."""
        order = {art: index for index, art in enumerate(self.export.article_numbers)}

        def key(record):
            parts = [order.get(record.art_no, len(order))]
            for name in keys:
                value = getattr(record, name)
                parts.append((value if isinstance(value, int) else _to_int(str(value))) if name == "sort_no" else str(value))
            return tuple(parts)

        return sorted(records, key=key)

    def genart_name(self, genart_id: str) -> str:
        name = self.genart_names.get(genart_id, "")
        if not name and genart_id and genart_id not in self.missing_genarts:
            self.missing_genarts.add(genart_id)
            self.report.add_hint("GenArt-Katalog", "", f"GenArt {genart_id} nicht im Katalog - Bezeichnung leer")
        return name

    # -- Generischer Datei-Merge --------------------------------------------

    def merge_file(
        self,
        file_spec: tuple[str, str],
        headers: list[str],
        ims_rows_by_key: dict[str, list[list[str]]],
        *,
        skip_existing_silently: bool = False,
    ) -> set[str]:
        """Haengt IMS-Zeilen je Artikel an, sofern die GUI den Artikel in der Datei
        nicht kennt. Liefert die Artikel-Schluessel mit Konflikt."""
        file_name, sheet_name = file_spec
        path = self.output_dir / file_name
        stats = self.report.stats(file_name)
        existing_rows = m.prepare_existing_rows_for_write(path, sheet_name, headers) if path.exists() else []
        stats.kept = len(existing_rows)
        gui_rows_by_key: dict[str, list[list[str]]] = {}
        for row in existing_rows:
            key = m.normalize_article_number(str(row[0])) if row else ""
            if key:
                gui_rows_by_key.setdefault(key, []).append(row)

        new_rows: list[list[str]] = []
        conflict_keys: set[str] = set()
        for key, rows in ims_rows_by_key.items():
            if not rows:
                continue
            stats.ims_rows += len(rows)
            if key in gui_rows_by_key:
                conflict_keys.add(key)
                stats.discarded += len(rows)
                if not skip_existing_silently:
                    stats.conflict_articles += 1
                    gui_cells = [[str(cell) for cell in row[:-1]] for row in gui_rows_by_key[key]]
                    ims_cells = [[str(cell) for cell in row] for row in rows]
                    identical = gui_cells == ims_cells
                    if not identical and file_name in ORDER_INSENSITIVE_FILES:
                        # GenArts sortiert die App ohnehin; bei Videos/Web Links ist die Reihenfolge ohne Bedeutung.
                        identical = sorted(gui_cells) == sorted(ims_cells)
                    if identical:
                        stats.identical_conflicts += 1
                    self.report.conflicts.append(
                        Conflict(
                            article=rows[0][0],
                            file_name=file_name,
                            gui_state=_compact_rows(gui_rows_by_key[key], drop_last=True),
                            ims_state=_compact_rows(rows, drop_last=False),
                            identical=identical,
                        )
                    )
                continue
            for row in rows:
                new_rows.append([str(cell) for cell in row] + [self.written_at])
        if new_rows:
            m.write_workbook(path, sheet_name, headers, existing_rows + new_rows)
            stats.written = True
        stats.appended = len(new_rows)
        return conflict_keys

    def gui_keys_in_file(self, file_spec: tuple[str, str], headers: list[str]) -> set[str]:
        path = self.output_dir / file_spec[0]
        if not path.exists():
            return set()
        return {
            m.normalize_article_number(str(row[0]))
            for row in m.read_workbook_rows(path, file_spec[1], len(headers))
            if row and m.normalize_article_number(str(row[0]))
        }

    # -- Texte ----------------------------------------------------------------

    def merge_short_texts(self) -> None:
        registry = m.IdRegistry()
        registry.load_from_folder(self.output_dir)
        for txt_mod_no in self.export.text_modules:
            registry.reserve(txt_mod_no)
        for info in self.export.info_texts:
            registry.reserve(info.txt_mod_no)
        ims_rows: dict[str, list[list[str]]] = {}
        mapping_rows: dict[str, list[list[str]]] = {}
        gui_short_keys = self.gui_keys_in_file(m.SHORT_TEXT_FILE, m.SHORT_TEXT_HEADERS)
        for article in self.export.articles:
            if not article.term_no:
                self.report.add_hint("Kurzbezeichnung", article.art_no, "Artikel ohne Kurzbezeichnung (200.TermNo leer)")
                continue
            translations = short_text_translations(self.export, article.term_no)
            if not any(translations.values()):
                self.report.add_hint("Kurzbezeichnung", article.art_no, f"TermNo {article.term_no} ohne 030-Saetze")
                continue
            key = self.key_of[article.art_no]
            if key in gui_short_keys:
                # Konflikt: keine neue ID verbrauchen, der Merge meldet ihn.
                module_id = registry.get_ids_for_article(article.art_no)[0]
                ims_rows.setdefault(key, []).append(build_text_row(article.art_no, module_id, "1", translations))
                continue
            module_id = registry.get_ids_for_article(article.art_no)[0]
            if not module_id:
                module_id = registry.generate_unique()
                registry.remember_article_ids(article.art_no, module_id, "")
            ims_rows.setdefault(key, []).append(build_text_row(article.art_no, module_id, "1", translations))
            mapping_rows.setdefault(key, []).append([article.art_no, module_id])
        conflicts = self.merge_file(m.SHORT_TEXT_FILE, m.SHORT_TEXT_HEADERS, ims_rows)
        mapping_rows = {key: rows for key, rows in mapping_rows.items() if key not in conflicts}
        self.merge_file(m.SHORT_MAPPING_FILE, m.SHORT_MAPPING_HEADERS, mapping_rows, skip_existing_silently=True)

    def merge_long_texts(self) -> None:
        ims_rows: dict[str, list[list[str]]] = {}
        seen_articles: set[str] = set()
        for info in self.export.info_texts:
            if info.sort_no != "01":
                self.report.add_hint(
                    "Langtext",
                    info.art_no,
                    f"206 SortNo {info.sort_no} mit Modul {info.txt_mod_no} nicht uebernommen (nur SortNo 01)",
                )
                continue
            if info.art_no in seen_articles:
                self.report.add_hint("Langtext", info.art_no, f"Zweiter 206-Satz mit SortNo 01 (Modul {info.txt_mod_no}) uebersprungen")
                continue
            if info.txt_mod_no not in self.export.text_modules:
                self.report.add_hint("Langtext", info.art_no, f"Textbaustein {info.txt_mod_no} fehlt in 035")
                continue
            seen_articles.add(info.art_no)
            translations = long_text_translations(self.export, info.txt_mod_no)
            key = self.key_of.get(info.art_no, m.normalize_article_number(info.art_no))
            ims_rows.setdefault(key, []).append(build_text_row(info.art_no, info.txt_mod_no, "2", translations))
        self.merge_file(m.LONG_TEXT_FILE, m.SHORT_TEXT_HEADERS, ims_rows)

    # -- GenArt / Attribute / Referenzen --------------------------------------

    def merge_genarts(self) -> None:
        ims_rows: dict[str, list[list[str]]] = {}
        for genart in self.export.genarts:
            genart_id = strip_leading_zeros(genart.genart_no)
            key = self.key_of.get(genart.art_no, m.normalize_article_number(genart.art_no))
            ims_rows.setdefault(key, []).append([genart.art_no, genart_id, self.genart_name(genart_id)])
        self.merge_file(m.GENART_FILE, m.GENART_HEADERS, ims_rows)

    def merge_attributes(self) -> None:
        ims_rows: dict[str, list[list[str]]] = {}
        for criterion in self._ordered(self.export.criteria, "sort_no"):
            criteria_id = strip_leading_zeros(criterion.crit_no)
            option = self.attribute_options.get(criteria_id)
            label = option.label if option is not None else ""
            value_format = option.value_format if option is not None else ""
            if option is None and criteria_id not in self.missing_criteria:
                self.missing_criteria.add(criteria_id)
                self.report.add_hint("Attribut-Katalog", criterion.art_no, f"Kriterium {criteria_id} nicht im Katalog - Bezeichnung/Format leer")
            value, value_to = criterion.crit_val, ""
            if value_format.casefold().startswith("numerisch"):
                value, value_to = split_numeric_interval(criterion.crit_val)
            key = self.key_of.get(criterion.art_no, m.normalize_article_number(criterion.art_no))
            ims_rows.setdefault(key, []).append([criterion.art_no, criteria_id, label, value_format, value, value_to])
        self.merge_file(m.ATTRIBUTE_FILE, m.ATTRIBUTE_HEADERS, ims_rows)

    def merge_trade_numbers(self) -> None:
        ims_rows: dict[str, list[list[str]]] = {}
        for trade in self._ordered(self.export.trade_numbers, "sort_no"):
            key = self.key_of.get(trade.art_no, m.normalize_article_number(trade.art_no))
            ims_rows.setdefault(key, []).append([trade.art_no, m.USAGE_NUMBER_TECDOC_ID, trade.trade_no])
        # Hinweis: GUI-Suchwoerter ueber 35 Zeichen passen nicht in 207.TradeNo.
        path = self.output_dir / m.USAGE_NUMBER_FILE[0]
        if path.exists():
            long_terms: dict[str, int] = {}
            for row in m.read_workbook_rows(path, m.USAGE_NUMBER_FILE[1], len(m.USAGE_NUMBER_HEADERS)):
                if len(str(row[2]).strip()) > 35:
                    long_terms[str(row[0])] = long_terms.get(str(row[0]), 0) + 1
            for article, count in long_terms.items():
                self.report.add_hint("Suchwort > 35 Zeichen", article, f"{count} GUI-Suchwort/-woerter laenger als 35 Zeichen (passen nicht in TAF 207)")
        self.merge_file(m.USAGE_NUMBER_FILE, m.USAGE_NUMBER_HEADERS, ims_rows)

    def merge_reference_numbers(self) -> None:
        oe_rows: dict[str, list[list[str]]] = {}
        comparison_rows: dict[str, list[list[str]]] = {}
        unknown_manufacturers: dict[str, int] = {}
        for reference in self._ordered(self.export.reference_numbers, "man_no", "sort_no"):
            manufacturer_id = strip_leading_zeros(reference.man_no)
            key = self.key_of.get(reference.art_no, m.normalize_article_number(reference.art_no))
            if manufacturer_id in self.competitors:
                comparison_rows.setdefault(key, []).append([reference.art_no, "10", manufacturer_id, reference.ref_no])
                continue
            if manufacturer_id not in self.manufacturers:
                unknown_manufacturers[reference.man_no] = unknown_manufacturers.get(reference.man_no, 0) + 1
            oe_rows.setdefault(key, []).append([reference.art_no, "4", manufacturer_id, reference.ref_no])
        for man_no, count in unknown_manufacturers.items():
            self.report.add_hint(
                "Hersteller-Katalog",
                "",
                f"ManNo {man_no} nicht in KHer.csv - {count} Satz/Saetze als OE-Nummer mit KHerNr {strip_leading_zeros(man_no)} uebernommen",
            )
        self.merge_file(m.OE_FILE, m.OE_HEADERS, oe_rows)
        self.merge_file(m.COMPARISON_FILE, m.COMPARISON_HEADERS, comparison_rows)

    # -- Fahrzeuge / Artikellisten ---------------------------------------------

    def merge_vehicle_links(self) -> None:
        ims_rows: dict[str, list[list[str]]] = {}
        topmotive_by_article: dict[str, int] = {}
        unknown_by_article: dict[str, int] = {}
        for link in self.export.linkages:
            genart_id = strip_leading_zeros(link.genart_no)
            ktyp = strip_leading_zeros(link.target_no)
            key = self.key_of.get(link.art_no, m.normalize_article_number(link.art_no))
            ims_rows.setdefault(key, []).append(
                [link.art_no, strip_leading_zeros(link.target_type), ktyp, "TecDoc", genart_id, self.genart_name(genart_id)]
            )
            if self.vehicle_catalog is not None and ktyp not in self.vehicle_catalog.by_ktyp_tecdoc:
                if ktyp in self.vehicle_catalog.by_ktyp_topmotive:
                    topmotive_by_article[link.art_no] = topmotive_by_article.get(link.art_no, 0) + 1
                else:
                    unknown_by_article[link.art_no] = unknown_by_article.get(link.art_no, 0) + 1
        for article, count in topmotive_by_article.items():
            self.report.add_hint("Topmotive-KTyp", article, f"{count} 400-Saetze mit Topmotive-Nummer statt TecDoc-KTypNr (im IMS so vorhanden)")
        for article, count in unknown_by_article.items():
            self.report.add_hint("KTyp unbekannt", article, f"{count} 400-Saetze mit KTypNr ausserhalb KTyp.xlsx")
        self.merge_file(m.VEHICLE_LINK_FILE, m.VEHICLE_LINK_HEADERS, ims_rows)

    def merge_article_lists(self) -> None:
        rows_by_article: dict[str, list[list[str]]] = {}
        for part in sorted(self.export.parts_lists, key=lambda item: (item.art_no, _to_int(item.seq_no))):
            genart_id = strip_leading_zeros(part.part_genart_no)
            rows_by_article.setdefault(part.art_no, []).append(
                [
                    part.art_no,
                    m.ARTICLE_LIST_TYPES[0],
                    part.part_no,
                    str(_to_int(part.quantity)),
                    genart_id,
                    self.genart_name(genart_id),
                    str(_to_int(part.seq_no)),
                ]
            )
        # 222: IMS expandiert je GenArt des Zubehoerartikels -> je (SeqNo, AccArtNo) eine Zeile,
        # erste GenArt, Sort Nr fortlaufend in Reihenfolge des ersten Auftretens.
        first_seen: dict[tuple[str, str, str], list[str]] = {}
        order: list[tuple[str, str, str]] = []
        for accessory in self.export.accessory_lists:
            key = (accessory.art_no, accessory.seq_no, accessory.acc_art_no)
            if key in first_seen:
                continue
            genart_id = strip_leading_zeros(accessory.acc_genart_no)
            first_seen[key] = [
                accessory.art_no,
                m.ARTICLE_LIST_TYPES[1],
                accessory.acc_art_no,
                str(_to_int(accessory.quantity)),
                genart_id,
                self.genart_name(genart_id),
                "",
            ]
            order.append(key)
        sort_numbers: dict[str, int] = {}
        for key in order:
            row = first_seen[key]
            sort_numbers[row[0]] = sort_numbers.get(row[0], 0) + 1
            row[6] = str(sort_numbers[row[0]])
            rows_by_article.setdefault(row[0], []).append(row)
        ims_rows: dict[str, list[list[str]]] = {}
        for art_no, rows in rows_by_article.items():
            key = self.key_of.get(art_no, m.normalize_article_number(art_no))
            ims_rows.setdefault(key, []).extend(rows)
        self.merge_file(m.ARTICLE_LIST_FILE, m.ARTICLE_LIST_HEADERS, ims_rows)

    # -- Medien ---------------------------------------------------------------

    def merge_media(self) -> None:
        allocations_by_article: dict[str, list[TafAllocation]] = {}
        for allocation in self.export.allocations:
            allocations_by_article.setdefault(allocation.art_no, []).append(allocation)
        for rows in allocations_by_article.values():
            rows.sort(key=lambda item: item.sort_no)

        gui_image_keys = self.gui_keys_in_file(m.IMAGE_FILE, m.IMAGE_HEADERS)
        gui_document_keys = self.gui_keys_in_file(m.DOCUMENT_FILE, m.DOCUMENT_HEADERS)

        # Vorpass: welche Artikel bekommen Dateien (media_limit zaehlt Artikel mit
        # tatsaechlich zu schreibenden Bild-/Dokumentdateien).
        file_articles: list[str] = []
        file_articles_set: set[str] = set()
        zip_refcount: dict[str, int] = {}
        if self.media:
            for article in self.export.articles:
                key = self.key_of[article.art_no]
                names: list[str] = []
                for allocation in allocations_by_article.get(article.art_no, []):
                    if allocation.doc_type in IMAGE_DOC_TYPES and key in gui_image_keys:
                        continue
                    if allocation.doc_type in DOCUMENT_DOC_TYPES and key in gui_document_keys:
                        continue
                    if allocation.doc_type in DOC_TYPE_ZIP_SUFFIX:
                        document = self.export.documents.get((allocation.doc_no, allocation.doc_type))
                        if document is not None:
                            names.append(document.doc_name + DOC_TYPE_ZIP_SUFFIX[allocation.doc_type])
                if not names:
                    continue
                if self.media_limit is not None and len(file_articles) >= self.media_limit:
                    break
                file_articles.append(article.art_no)
                file_articles_set.add(article.art_no)
                for name in names:
                    zip_refcount[name] = zip_refcount.get(name, 0) + 1

        image_rows: dict[str, list[list[str]]] = {}
        document_rows: dict[str, list[list[str]]] = {}
        video_rows: dict[str, list[list[str]]] = {}
        web_rows: dict[str, list[list[str]]] = {}
        pic_zip = zipfile.ZipFile(self.export.pic_zip) if self.media and self.export.pic_zip else None
        doc_zip = zipfile.ZipFile(self.export.doc_zip) if self.media and self.export.doc_zip else None
        zip_names = {}
        if pic_zip is not None:
            zip_names.update({name: pic_zip for name in pic_zip.namelist()})
        if doc_zip is not None:
            zip_names.update({name: doc_zip for name in doc_zip.namelist()})
        byte_cache: dict[str, bytes] = {}

        def release_member(name: str) -> None:
            # Referenzzaehler des Zip-Caches: geteilte Dateien werden einmal
            # gelesen und nach dem letzten Artikel aus dem Speicher entlassen.
            zip_refcount[name] = zip_refcount.get(name, 1) - 1
            if zip_refcount[name] <= 0:
                byte_cache.pop(name, None)

        def read_member(name: str) -> bytes | None:
            data = byte_cache.get(name)
            if data is None:
                archive = zip_names.get(name)
                if archive is None:
                    release_member(name)
                    return None
                data = archive.read(name)
                if zip_refcount.get(name, 0) > 1:
                    byte_cache[name] = data
            release_member(name)
            return data

        missing_members: set[str] = set()
        try:
            for article in self.export.articles:
                art_no = article.art_no
                key = self.key_of[art_no]
                allocations = allocations_by_article.get(art_no)
                if not allocations:
                    continue
                with_files = art_no in file_articles_set
                target_dir: Path | None = None
                index: dict[str, str] = {}
                index_changed = False
                for allocation in allocations:
                    document = self.export.documents.get((allocation.doc_no, allocation.doc_type))
                    if document is None:
                        self.report.add_hint("Medien", art_no, f"232 verweist auf fehlenden 231-Satz {allocation.doc_no}/{allocation.doc_type}")
                        continue
                    if allocation.doc_type == URL_DOC_TYPE:
                        kind, link = classify_url(document.url)
                        if kind == "video":
                            video_rows.setdefault(key, []).append([art_no, link, "4"])
                        else:
                            web_rows.setdefault(key, []).append([art_no, link, "4"])
                        continue
                    if allocation.doc_type not in DOC_TYPE_ZIP_SUFFIX:
                        self.report.add_hint("Medien", art_no, f"DocType {allocation.doc_type} ({document.doc_name}) nicht unterstuetzt")
                        continue
                    if not with_files:
                        continue
                    is_image = allocation.doc_type in IMAGE_DOC_TYPES
                    if is_image and key in gui_image_keys:
                        continue
                    if not is_image and key in gui_document_keys:
                        continue
                    if target_dir is None:
                        target_dir = media_target_dir(self.output_dir, art_no)
                        index = m.load_media_index(target_dir)
                    member_name = document.doc_name + DOC_TYPE_ZIP_SUFFIX[allocation.doc_type]
                    source_key = f"{MEDIA_SOURCE_PREFIX}{allocation.doc_no}/{allocation.doc_type}"
                    suffix = DOC_TYPE_GUI_SUFFIX[allocation.doc_type]
                    existing_name = next(
                        (name for name, source in index.items() if source == source_key and name.lower().endswith(suffix)),
                        None,
                    )
                    if existing_name is not None and m.path_exists_safe(target_dir / existing_name) and (target_dir / existing_name).stat().st_size > 0:
                        target_path = target_dir / existing_name
                        release_member(member_name)
                        self.report.media_files_reused += 1
                    else:
                        data = read_member(member_name)
                        if data is None:
                            if member_name not in missing_members:
                                missing_members.add(member_name)
                            self.report.add_hint("Medien", art_no, f"Datei {member_name} fehlt im Zip - keine Zeile geschrieben")
                            continue
                        target_dir.mkdir(parents=True, exist_ok=True)
                        stem = m._allocate_media_stem(target_dir, art_no, "" if is_image else "D")
                        target_path = target_dir / f"{stem}{suffix}"
                        if m.path_exists_safe(target_path):
                            self.report.add_hint("Medien", art_no, f"{target_path.name} existiert bereits - nicht ueberschrieben")
                            continue
                        target_path.write_bytes(data)
                        index[target_path.name] = source_key
                        index_changed = True
                        self.report.media_files_written += 1
                    if is_image:
                        image_rows.setdefault(key, []).append(
                            [art_no, str(target_path), "5", "255", strip_leading_zeros(allocation.doc_type)]
                        )
                    else:
                        document_rows.setdefault(key, []).append(
                            [art_no, str(target_path), "255", strip_leading_zeros(document.term_norm) or "17", strip_leading_zeros(allocation.doc_type)]
                        )
                if index_changed and target_dir is not None:
                    m.save_media_index(target_dir, index)
                if with_files:
                    self.report.media_articles += 1

            # Markenlogo (DocContentType 002) ist kein Artikelmedium.
            if self.media:
                for document in self.export.documents.values():
                    if document.content_type != LOGO_CONTENT_TYPE or document.doc_type not in DOC_TYPE_ZIP_SUFFIX:
                        continue
                    member_name = document.doc_name + DOC_TYPE_ZIP_SUFFIX[document.doc_type]
                    archive = zip_names.get(member_name)
                    if archive is None:
                        self.report.add_hint("Logo", "", f"Logo {member_name} fehlt im Zip")
                        continue
                    logo_dir = (self.output_dir / m.MEDIA_SUBFOLDER / LOGO_SUBFOLDER).resolve()
                    logo_dir.mkdir(parents=True, exist_ok=True)
                    logo_path = logo_dir / member_name
                    if not m.path_exists_safe(logo_path):
                        logo_path.write_bytes(archive.read(member_name))
                    self.report.logo_path = str(logo_path)
        finally:
            if pic_zip is not None:
                pic_zip.close()
            if doc_zip is not None:
                doc_zip.close()

        self.merge_file(m.IMAGE_FILE, m.IMAGE_HEADERS, image_rows)
        self.merge_file(m.DOCUMENT_FILE, m.DOCUMENT_HEADERS, document_rows)
        self.merge_file(m.VIDEO_FILE, m.VIDEO_HEADERS, video_rows)
        self.merge_file(m.WEB_LINK_FILE, m.WEB_HEADERS, web_rows)

    # -- Nur zaehlen ---------------------------------------------------------------

    def repair_empty_module_ids(self) -> None:
        """GUI-Zeilen ohne Text-Modul-ID (Altbestand) bekommen eine ID: Langtexte die
        Modul-ID aus Apollo (206), damit Apollo kein zweites Modul anlegt; Kurz-
        bezeichnungen eine neue ID der App (dazu die Zuordnungszeile)."""
        long_ids = {info.art_no: info.txt_mod_no for info in self.export.info_texts if info.sort_no == "01"}
        long_by_key = {m.normalize_article_number(art): mod for art, mod in long_ids.items()}
        path = self.output_dir / m.LONG_TEXT_FILE[0]
        if path.exists():
            rows = m.prepare_existing_rows_for_write(path, m.LONG_TEXT_FILE[1], m.SHORT_TEXT_HEADERS)
            changed = 0
            for row in rows:
                if row and not str(row[1]).strip():
                    module_id = long_by_key.get(m.normalize_article_number(str(row[0])), "")
                    if module_id:
                        row[1] = module_id
                        changed += 1
                        self.report.add_hint("Text-Modul-ID ergaenzt", str(row[0]), f"Langtext ohne ID -> Modul {module_id} aus Apollo (206)")
            if changed:
                m.write_workbook(path, m.LONG_TEXT_FILE[1], m.SHORT_TEXT_HEADERS, rows)
        path = self.output_dir / m.SHORT_TEXT_FILE[0]
        if path.exists():
            registry = m.IdRegistry()
            registry.load_from_folder(self.output_dir)
            for txt_mod_no in self.export.text_modules:
                registry.reserve(txt_mod_no)
            rows = m.prepare_existing_rows_for_write(path, m.SHORT_TEXT_FILE[1], m.SHORT_TEXT_HEADERS)
            mapping_path = self.output_dir / m.SHORT_MAPPING_FILE[0]
            mapping_rows = m.prepare_existing_rows_for_write(mapping_path, m.SHORT_MAPPING_FILE[1], m.SHORT_MAPPING_HEADERS) if mapping_path.exists() else []
            mapped = {m.normalize_article_number(str(row[0])) for row in mapping_rows if row}
            changed = 0
            for row in rows:
                if row and not str(row[1]).strip():
                    module_id = registry.get_ids_for_article(str(row[0]))[0] or registry.generate_unique()
                    registry.remember_article_ids(str(row[0]), module_id, "")
                    row[1] = module_id
                    changed += 1
                    if m.normalize_article_number(str(row[0])) not in mapped:
                        mapping_rows.append([str(row[0]), module_id, self.written_at])
                    self.report.add_hint("Text-Modul-ID ergaenzt", str(row[0]), f"Kurzbezeichnung ohne ID -> neue ID {module_id}")
            if changed:
                m.write_workbook(path, m.SHORT_TEXT_FILE[1], m.SHORT_TEXT_HEADERS, rows)
                m.write_workbook(mapping_path, m.SHORT_MAPPING_FILE[1], m.SHORT_MAPPING_HEADERS, mapping_rows)

    def count_only(self) -> None:
        self.report.counted_only["212.dat"] = len(self.export.statuses)
        self.report.counted_only["299.dat"] = len(self.export.guids)
        self.report.counted_only["209.dat (EAN, kein GUI-Ziel)"] = len(self.export.gtins)
        # Logo-Satz und Zubehoer-Kennzeichen haben kein GUI-Ziel.
        accessory_flags = sum(1 for article in self.export.articles if article.accessory)
        if accessory_flags:
            self.report.counted_only["200.Accessory=1"] = accessory_flags

    def run(self) -> TafSyncReport:
        started = time.perf_counter()
        self.merge_short_texts()
        self.merge_long_texts()
        self.merge_genarts()
        self.merge_attributes()
        self.merge_trade_numbers()
        self.merge_reference_numbers()
        self.merge_media()
        self.merge_vehicle_links()
        self.merge_article_lists()
        self.repair_empty_module_ids()
        self.count_only()
        self.report.duration_seconds = time.perf_counter() - started
        return self.report


def merge_taf_into_output(
    export: TafExport,
    output_dir: Path | str,
    *,
    media: bool = True,
    media_limit: int | None = None,
    written_at: str | None = None,
    report_path: Path | str | None = None,
    genart_catalog_path: Path | str = m.DEFAULT_GENART_SOURCE,
    attribute_catalog_path: Path | str = m.DEFAULT_ATTRIBUTE_SOURCE,
    competitor_catalog_path: Path | str = m.DEFAULT_COMPETITOR_SOURCE,
    vehicle_catalog_path: Path | str | None = m.DEFAULT_VEHICLE_SOURCE,
) -> TafSyncReport:
    """Uebernimmt den Export in den Output-Ordner der GUI und liefert den Bericht.

    media=False: keine Bild-/Dokumentdateien und -zeilen (Videos/Web Links werden
    immer geschrieben, sie zeigen auf keine Datei). media_limit=N: nur die ersten
    N Artikel (200-Reihenfolge) mit zu schreibenden Dateien bekommen Medienzeilen
    und Dateien, die uebrigen gar keine Bild-/Dokumentzeilen.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    merger = _Merger(
        export,
        output_dir,
        media=media,
        media_limit=media_limit,
        written_at=written_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        genart_catalog_path=Path(genart_catalog_path),
        attribute_catalog_path=Path(attribute_catalog_path),
        competitor_catalog_path=Path(competitor_catalog_path),
        vehicle_catalog_path=Path(vehicle_catalog_path) if vehicle_catalog_path else None,
    )
    try:
        report = merger.run()
    except Exception:
        # Teilbericht retten: welche Dateien bis zum Abbruch ergaenzt wurden und
        # welche Konflikte/Hinweise anfielen, stuende sonst nirgends.
        if report_path is not None:
            try:
                merger.report.write_xlsx(Path(report_path))
            except Exception:
                pass
        raise
    if report_path is not None:
        report.write_xlsx(Path(report_path))
    return report


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Apollo-IMS-TAF-2.7-Export in die Output-Dateien der Import-GUI uebernehmen.")
    parser.add_argument("export_dir", help="Exportordner mit 0757\\*.dat, 0757_PIC.zip, 0757_DOC.zip")
    parser.add_argument("output_dir", help="Output-Ordner der Import-GUI (Kopie!)")
    parser.add_argument("--no-media", action="store_true", help="keine Bild-/Dokumentdateien und -zeilen schreiben")
    parser.add_argument("--media-limit", type=int, default=None, metavar="N", help="nur die ersten N Artikel mit Medien")
    parser.add_argument("--report", default=None, metavar="XLSX", help="Bericht als Excel-Datei")
    parser.add_argument("--written-at", default=None, help="Zeitstempel fuer 'Zuletzt geschrieben am' (Standard: jetzt)")
    args = parser.parse_args(argv)

    started = time.perf_counter()
    export = read_taf_export(Path(args.export_dir))
    print(f"Export gelesen: {len(export.articles)} Artikel, " + ", ".join(f"{k}={v}" for k, v in export.record_counts.items()))
    report = merge_taf_into_output(
        export,
        Path(args.output_dir),
        media=not args.no_media,
        media_limit=args.media_limit,
        written_at=args.written_at,
        report_path=Path(args.report) if args.report else None,
    )
    for line in report.summary_lines():
        print(line)
    print(f"Gesamtlaufzeit {time.perf_counter() - started:.1f} s")
    return 0


if __name__ == "__main__":
    sys.exit(main())
